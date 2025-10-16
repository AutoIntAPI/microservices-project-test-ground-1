#!/usr/bin/env python3
"""
API Call Extractor for Microservices Project

This script extracts all REST API calls from the microservices codebase,
including API Gateway routes and inter-service communication.
"""

import os
import re
import json
import csv
from pathlib import Path
from typing import List, Dict, Tuple

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

# Define service directories
SERVICE_DIRS = {
    'api-gateway': 'services/api-gateway/src',
    'user-service': 'services/user-service/src',
    'product-service': 'services/product-service/src',
    'order-service': 'services/order-service/app',
    'payment-service': 'services/payment-service/app',
    'notification-service': 'services/notification-service/app',
}

# HTTP methods
HTTP_METHODS = ['GET', 'POST', 'PUT', 'DELETE', 'PATCH', 'OPTIONS', 'HEAD']


class APICallExtractor:
    def __init__(self, repo_path: str):
        self.repo_path = Path(repo_path)
        self.api_calls = []
        
    def extract_gateway_routes(self):
        """Extract API Gateway routes that proxy to backend services"""
        gateway_dir = self.repo_path / SERVICE_DIRS['api-gateway'] / 'routes'
        
        if not gateway_dir.exists():
            print(f"Warning: Gateway directory not found: {gateway_dir}")
            return
        
        for route_file in gateway_dir.glob('*.js'):
            service_name = route_file.stem  # e.g., 'users', 'products'
            content = route_file.read_text()
            
            # Extract service URL
            service_url_match = re.search(
                r'const\s+\w+_SERVICE_URL\s*=\s*.*?\|\|\s*[\'"]([^\'"]+)[\'"]',
                content
            )
            destination_url = service_url_match.group(1) if service_url_match else 'unknown'
            destination_service = destination_url.split('://')[1].split(':')[0] if '://' in destination_url else 'unknown'
            
            # Extract routes
            # Pattern: router.METHOD('PATH', ...)
            route_pattern = r'router\.(get|post|put|delete|patch)\s*\(\s*[\'"]([^\'"]+)[\'"]'
            
            for match in re.finditer(route_pattern, content, re.IGNORECASE):
                method = match.group(1).upper()
                endpoint = match.group(2)
                
                # Check if authentication is required
                auth_required = self._check_auth_in_route(content, endpoint)
                
                # Construct full API Gateway path
                gateway_path = f'/{service_name}{endpoint}'
                
                self.api_calls.append({
                    'source_service': 'api-gateway',
                    'source_file': str(route_file.relative_to(self.repo_path)),
                    'http_method': method,
                    'endpoint': gateway_path,
                    'destination_service': destination_service,
                    'destination_url': destination_url,
                    'authentication_required': auth_required,
                    'call_type': 'proxy'
                })
    
    def extract_nodejs_service_calls(self, service_name: str):
        """Extract REST API calls from Node.js services"""
        service_dir = self.repo_path / SERVICE_DIRS[service_name]
        
        if not service_dir.exists():
            print(f"Warning: Service directory not found: {service_dir}")
            return
        
        for js_file in service_dir.rglob('*.js'):
            content = js_file.read_text()
            
            # Look for axios or fetch or http/https requests
            # Pattern: axios.get/post/etc('url') or fetch('url')
            request_patterns = [
                r'axios\.(get|post|put|delete|patch)\s*\(\s*[\'"`]([^\'"` ]+)[\'"`]',
                r'fetch\s*\(\s*[\'"`]([^\'"` ]+)[\'"`].*?method:\s*[\'"`](\w+)[\'"`]',
                r'requests?\.(get|post|put|delete|patch)\s*\(\s*[\'"`]([^\'"` ]+)[\'"`]',
            ]
            
            for pattern in request_patterns:
                for match in re.finditer(pattern, content, re.IGNORECASE):
                    if 'axios' in pattern or 'requests' in pattern:
                        method = match.group(1).upper()
                        url = match.group(2)
                    else:  # fetch
                        url = match.group(1)
                        method = match.group(2).upper() if len(match.groups()) > 1 else 'GET'
                    
                    # Extract destination service from URL
                    destination_service, destination_url = self._parse_destination(url, content)
                    
                    self.api_calls.append({
                        'source_service': service_name,
                        'source_file': str(js_file.relative_to(self.repo_path)),
                        'http_method': method,
                        'endpoint': url,
                        'destination_service': destination_service,
                        'destination_url': destination_url,
                        'authentication_required': 'unknown',
                        'call_type': 'rest_call'
                    })
    
    def extract_python_service_calls(self, service_name: str):
        """Extract REST API calls from Python/Flask services"""
        service_dir = self.repo_path / SERVICE_DIRS[service_name]
        
        if not service_dir.exists():
            print(f"Warning: Service directory not found: {service_dir}")
            return
        
        for py_file in service_dir.rglob('*.py'):
            content = py_file.read_text()
            
            # Look for requests library calls - handle multi-line calls
            # Pattern: requests.get/post/etc(...)
            request_pattern = r'(?:response\s*=\s*)?requests\.(get|post|put|delete|patch)\s*\(\s*f?[\'"]([^\'"]+)[\'"]'
            
            for match in re.finditer(request_pattern, content, re.IGNORECASE | re.MULTILINE):
                method = match.group(1).upper()
                url_pattern = match.group(2)
                
                # Extract destination service from URL or environment variable
                destination_service, destination_url = self._parse_destination_python(url_pattern, content)
                
                self.api_calls.append({
                    'source_service': service_name,
                    'source_file': str(py_file.relative_to(self.repo_path)),
                    'http_method': method,
                    'endpoint': url_pattern,
                    'destination_service': destination_service,
                    'destination_url': destination_url,
                    'authentication_required': 'unknown',
                    'call_type': 'rest_call'
                })
    
    def extract_service_endpoints(self, service_name: str):
        """Extract service endpoints that are exposed"""
        service_dir = self.repo_path / SERVICE_DIRS[service_name]
        
        if not service_dir.exists():
            return
        
        # Track already added endpoints to avoid duplicates
        seen_endpoints = set()
        
        # For Node.js services, look in routes directory
        routes_dir = service_dir / 'routes'
        if routes_dir.exists():
            for route_file in routes_dir.glob('*.js'):
                content = route_file.read_text()
                
                # Pattern: router.METHOD('PATH', ...)
                route_pattern = r'router\.(get|post|put|delete|patch)\s*\(\s*[\'"]([^\'"]+)[\'"]'
                
                for match in re.finditer(route_pattern, content, re.IGNORECASE):
                    method = match.group(1).upper()
                    endpoint = match.group(2)
                    
                    key = (service_name, method, endpoint)
                    if key not in seen_endpoints:
                        seen_endpoints.add(key)
                        self.api_calls.append({
                            'source_service': service_name,
                            'source_file': str(route_file.relative_to(self.repo_path)),
                            'http_method': method,
                            'endpoint': endpoint,
                            'destination_service': service_name,
                            'destination_url': f'http://{service_name}:PORT',
                            'authentication_required': 'varies',
                            'call_type': 'endpoint_definition'
                        })
        
        # For Python services, look for Flask blueprints
        for py_file in service_dir.rglob('*.py'):
            content = py_file.read_text()
            
            # Pattern: @bp.route('PATH', methods=['METHOD'])
            route_pattern = r'@bp\.route\s*\(\s*[\'"]([^\'"]+)[\'"].*?methods\s*=\s*\[([^\]]+)\]'
            
            for match in re.finditer(route_pattern, content, re.IGNORECASE | re.DOTALL):
                endpoint = match.group(1)
                methods_str = match.group(2)
                
                # Extract individual methods
                methods = re.findall(r'[\'"](\w+)[\'"]', methods_str)
                
                for method in methods:
                    key = (service_name, method.upper(), endpoint)
                    if key not in seen_endpoints:
                        seen_endpoints.add(key)
                        self.api_calls.append({
                            'source_service': service_name,
                            'source_file': str(py_file.relative_to(self.repo_path)),
                            'http_method': method.upper(),
                            'endpoint': endpoint,
                            'destination_service': service_name,
                            'destination_url': f'http://{service_name}:PORT',
                            'authentication_required': 'varies',
                            'call_type': 'endpoint_definition'
                        })
    
    def _check_auth_in_route(self, content: str, endpoint: str) -> str:
        """Check if authentication middleware is applied to a route"""
        # Look for authenticate middleware in the same line or before the route
        lines = content.split('\n')
        
        for i, line in enumerate(lines):
            if endpoint in line:
                # Check current line and a few lines before
                context = '\n'.join(lines[max(0, i-5):i+1])
                if 'authenticate' in context:
                    return 'yes'
                elif 'router.use(authenticate)' in content:
                    return 'yes'
        
        return 'no'
    
    def _parse_destination(self, url_pattern: str, file_content: str) -> Tuple[str, str]:
        """Parse destination service and URL from URL pattern"""
        # Check if it's a variable reference
        if '{' in url_pattern or '$' in url_pattern:
            # Extract variable name
            var_match = re.search(r'[{$]\{?(\w+)', url_pattern)
            if var_match:
                var_name = var_match.group(1)
                # Find variable definition in content
                var_pattern = rf'{var_name}\s*=.*?[\'"]([^\'"]+)[\'"]'
                var_def = re.search(var_pattern, file_content)
                if var_def:
                    base_url = var_def.group(1)
                    # Extract service name from URL
                    if '://' in base_url:
                        service = base_url.split('://')[1].split(':')[0]
                        return service, base_url
        
        # Direct URL
        if '://' in url_pattern:
            service = url_pattern.split('://')[1].split(':')[0].split('/')[0]
            return service, url_pattern
        
        # Relative URL - try to find service URL constant
        for service in ['PRODUCT_SERVICE_URL', 'USER_SERVICE_URL', 'ORDER_SERVICE_URL', 
                       'PAYMENT_SERVICE_URL', 'NOTIFICATION_SERVICE_URL']:
            if service in file_content:
                url_match = re.search(
                    rf'{service}\s*=.*?[\'"]([^\'"]+)[\'"]',
                    file_content
                )
                if url_match:
                    base_url = url_match.group(1)
                    if '://' in base_url:
                        dest_service = base_url.split('://')[1].split(':')[0]
                        return dest_service, base_url
        
        return 'unknown', 'unknown'
    
    def _parse_destination_python(self, url_pattern: str, file_content: str) -> Tuple[str, str]:
        """Parse destination service and URL from URL pattern in Python files"""
        original_pattern = url_pattern
        
        # Handle f-string variable references like {PRODUCT_SERVICE_URL}
        if '{' in url_pattern and '}' in url_pattern:
            # Extract variable name from f-string
            var_matches = re.findall(r'\{([^}]+)\}', url_pattern)
            for var_name in var_matches:
                # Clean variable name (remove any formatting)
                clean_var = var_name.split(':')[0].strip()
                
                # Special handling for common variable names
                if '_SERVICE_URL' in clean_var or 'SERVICE_URL' in clean_var:
                    # Find variable definition in content - handle os.getenv() pattern
                    var_pattern = rf'{clean_var}\s*=\s*os\.getenv\([^,]+,\s*[\'"]([^\'"]+)[\'"]'
                    var_def = re.search(var_pattern, file_content)
                    if not var_def:
                        # Try simpler pattern
                        var_pattern = rf'{clean_var}\s*=.*?[\'"]([^\'"]+)[\'"]'
                        var_def = re.search(var_pattern, file_content)
                    
                    if var_def:
                        base_url = var_def.group(1)
                        # Replace the variable in the URL pattern
                        url_pattern = url_pattern.replace(f'{{{var_name}}}', base_url)
        
        # Direct URL
        if '://' in url_pattern:
            parts = url_pattern.split('://')[1].split(':')
            service = parts[0]
            # Extract just the base URL
            if ':' in url_pattern.split('://')[1]:
                port_and_rest = url_pattern.split('://')[1].split(':')[1]
                port = port_and_rest.split('/')[0]
                base_url = f"http://{service}:{port}"
            else:
                base_url = f"http://{service}"
            return service, base_url
        
        # If still has variable reference, try to resolve from environment variables in file
        if '{' in original_pattern:
            var_name = re.search(r'\{([^}]+)\}', original_pattern)
            if var_name:
                clean_var = var_name.group(1).split(':')[0].strip()
                # Try to find the variable definition with os.getenv
                var_pattern = rf'{clean_var}\s*=\s*os\.getenv\([^,]+,\s*[\'"]([^\'"]+)[\'"]'
                var_def = re.search(var_pattern, file_content)
                if not var_def:
                    # Try simpler pattern
                    var_pattern = rf'{clean_var}\s*=.*?[\'"]([^\'"]+)[\'"]'
                    var_def = re.search(var_pattern, file_content)
                
                if var_def:
                    base_url = var_def.group(1)
                    if '://' in base_url:
                        service = base_url.split('://')[1].split(':')[0]
                        return service, base_url
        
        return 'unknown', 'unknown'
    
    def extract_all(self):
        """Extract all API calls from all services"""
        print("Extracting API Gateway routes...")
        self.extract_gateway_routes()
        
        print("Extracting Node.js service calls and endpoints...")
        for service in ['user-service', 'product-service']:
            print(f"  - Processing {service}")
            self.extract_nodejs_service_calls(service)
            self.extract_service_endpoints(service)
        
        print("Extracting Python service calls and endpoints...")
        for service in ['order-service', 'payment-service', 'notification-service']:
            print(f"  - Processing {service}")
            self.extract_python_service_calls(service)
            self.extract_service_endpoints(service)
        
        print(f"\nTotal API calls extracted: {len(self.api_calls)}")
    
    def export_to_csv(self, output_path: str):
        """Export extracted API calls to CSV file"""
        with open(output_path, 'w', newline='', encoding='utf-8') as csvfile:
            fieldnames = [
                'source_service',
                'source_file', 
                'http_method',
                'endpoint',
                'destination_service',
                'destination_url',
                'authentication_required',
                'call_type'
            ]
            
            writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
            writer.writeheader()
            
            for call in sorted(self.api_calls, key=lambda x: (x['source_service'], x['http_method'], x['endpoint'])):
                writer.writerow(call)
        
        print(f"Exported to CSV: {output_path}")
    
    def export_to_json(self, output_path: str):
        """Export extracted API calls to JSON file"""
        with open(output_path, 'w', encoding='utf-8') as jsonfile:
            json.dump(self.api_calls, jsonfile, indent=2)
        
        print(f"Exported to JSON: {output_path}")
    
    def export_to_excel(self, output_path: str):
        """Export extracted API calls to Excel file"""
        if not EXCEL_AVAILABLE:
            print("Excel export not available. Install openpyxl: pip install openpyxl")
            return
        
        wb = Workbook()
        ws = wb.active
        ws.title = "API Calls"
        
        # Define headers
        headers = [
            'Source Service',
            'Source File', 
            'HTTP Method',
            'Endpoint',
            'Destination Service',
            'Destination URL',
            'Authentication Required',
            'Call Type'
        ]
        
        # Style headers
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        # Write headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        
        # Write data
        for row_idx, call in enumerate(sorted(self.api_calls, 
                                             key=lambda x: (x['source_service'], 
                                                          x['http_method'], 
                                                          x['endpoint'])), 2):
            ws.cell(row=row_idx, column=1, value=call['source_service'])
            ws.cell(row=row_idx, column=2, value=call['source_file'])
            ws.cell(row=row_idx, column=3, value=call['http_method'])
            ws.cell(row=row_idx, column=4, value=call['endpoint'])
            ws.cell(row=row_idx, column=5, value=call['destination_service'])
            ws.cell(row=row_idx, column=6, value=call['destination_url'])
            ws.cell(row=row_idx, column=7, value=call['authentication_required'])
            ws.cell(row=row_idx, column=8, value=call['call_type'])
        
        # Auto-adjust column widths
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width
        
        # Save workbook
        wb.save(output_path)
        print(f"Exported to Excel: {output_path}")


def main():
    # Get repository root directory
    script_dir = Path(__file__).parent
    repo_path = script_dir.parent
    
    print(f"Repository path: {repo_path}")
    print("=" * 60)
    
    # Create extractor
    extractor = APICallExtractor(repo_path)
    
    # Extract all API calls
    extractor.extract_all()
    
    # Export results
    output_dir = repo_path / 'api_extraction_output'
    output_dir.mkdir(exist_ok=True)
    
    csv_output = output_dir / 'api_calls.csv'
    json_output = output_dir / 'api_calls.json'
    excel_output = output_dir / 'api_calls.xlsx'
    
    extractor.export_to_csv(csv_output)
    extractor.export_to_json(json_output)
    extractor.export_to_excel(excel_output)
    
    print("=" * 60)
    print("API extraction complete!")
    print(f"\nOutput files:")
    print(f"  - CSV: {csv_output}")
    print(f"  - JSON: {json_output}")
    if EXCEL_AVAILABLE:
        print(f"  - Excel: {excel_output}")


if __name__ == '__main__':
    main()
